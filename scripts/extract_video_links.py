"""Extract BBMP e-Khata video links from EKHATA.xlsx into a clean JSON list."""
import json
import re
from pathlib import Path
from openpyxl import load_workbook

SRC = Path(r"C:\Users\hp\Downloads\EKHATA.xlsx")

STOP = set("""
a an the of for in on at to and or but by from with into over under as is are was were be been being do does did
i me you we they he she it him her us them my your our their what whats how where why when which who whose
this that these those there here can could would should will may might must shall do does did has have had
about with into onto over under again further once please thanks thank ok okay yes no not so just
e ekhata khata bbmp epid pid
""".split())


def keywords(title: str) -> list[str]:
    t = title.lower()
    t = re.sub(r"[^a-z0-9\s]", " ", t)
    toks = [w for w in t.split() if w and w not in STOP and len(w) > 1]
    seen = []
    for w in toks:
        if w not in seen:
            seen.append(w)
    return seen


def main():
    wb = load_workbook(SRC, data_only=True)
    ws = wb["Sheet1"]
    rows = list(ws.iter_rows(values_only=True))

    videos = []
    for r in rows:
        cells = [("" if c is None else str(c)).strip() for c in r]
        if not any(cells):
            continue
        # Skip header rows
        if cells[0].upper() == "SL NO":
            continue
        sl, title, link_en, link_kn = (cells + ["", "", "", ""])[:4]
        if not title:
            continue
        videos.append({
            "title": title.strip(),
            "link_en": link_en.strip(),
            "link_kn": link_kn.strip(),
            "keywords": keywords(title),
        })

    out = Path("scripts/bbmp-videos.json")
    out.write_text(json.dumps(videos, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"wrote {len(videos)} videos to {out}")


if __name__ == "__main__":
    main()
